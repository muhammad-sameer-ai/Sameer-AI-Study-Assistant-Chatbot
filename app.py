import os
import io
import streamlit as st
from pypdf import PdfReader
from docx import Document
from pptx import Presentation
from PIL import Image
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.vectorstores import Chroma
from langchain_google_genai import GoogleGenerativeAIEmbeddings, ChatGoogleGenerativeAI
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

st.set_page_config(page_title="Multi-Format Study Assistant", layout="wide")

st.title("📚 AI Study Assistant & Question Paper Generator")
st.caption("Upload PDFs, Word files (.docx), PowerPoint presentations (.pptx), or Images (.png, .jpg)")

with st.sidebar:
    st.header("1. Setup & Upload")
    api_key = st.text_input("Enter Gemini API Key", type="password")
    uploaded_files = st.file_uploader(
        "Upload Study Materials", 
        type=["pdf", "docx", "pptx", "png", "jpg", "jpeg"], 
        accept_multiple_files=True
    )

def extract_text_from_files(files, llm_instance):
    combined_text = ""
    for file in files:
        file_ext = file.name.split(".")[-1].lower()
        
        if file_ext == "pdf":
            reader = PdfReader(file)
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    combined_text += extracted + "\n"
                    
        elif file_ext == "docx":
            doc = Document(file)
            for para in doc.paragraphs:
                if para.text.strip():
                    combined_text += para.text + "\n"
                    
        elif file_ext == "pptx":
            prs = Presentation(file)
            for slide in prs.slides:
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            if paragraph.text.strip():
                                combined_text += paragraph.text + "\n"
                                
        elif file_ext in ["png", "jpg", "jpeg"]:
            try:
                image = Image.open(file)
                vision_response = llm_instance.invoke([
                    "Extract all text, equations, tables, and key information from this image verbatim.", 
                    image
                ])
                if vision_response.content:
                    combined_text += f"\n[Content from Image {file.name}]:\n" + vision_response.content + "\n"
            except Exception as e:
                st.warning(f"Could not read image {file.name}: {str(e)}")
            
    return combined_text

def create_pdf_paper(paper_text):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()
    
    custom_style = ParagraphStyle(
        'PaperStyle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=11,
        leading=16,
        spaceAfter=10
    )
    
    story = []
    lines = paper_text.split('\n')
    for line in lines:
        if line.strip():
            formatted_line = line.replace('**', '<b>').replace('**', '</b>')
            story.append(Paragraph(formatted_line, custom_style))
        else:
            story.append(Spacer(1, 8))
            
    doc.build(story)
    buffer.seek(0)
    return buffer

if api_key and uploaded_files:
    os.environ["GOOGLE_API_KEY"] = api_key
    llm = ChatGoogleGenerativeAI(model="gemini-2.5-flash", temperature=0.2)

    @st.cache_resource(show_spinner="Processing uploaded documents and images...")
    def process_documents(files):
        raw_text = extract_text_from_files(files, llm)
        if not raw_text.strip():
            st.error("No readable text could be extracted from the uploaded files. Please check your files.")
            st.stop()

        text_splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=200)
        chunks = text_splitter.split_text(raw_text)
        
        # Updated to modern Gemini Embedding Model
        embeddings = GoogleGenerativeAIEmbeddings(model="models/text-embedding-004")
        vector_store = Chroma.from_texts(chunks, embedding=embeddings)
        return vector_store

    vector_store = process_documents(uploaded_files)
    retriever = vector_store.as_retriever(search_kwargs={"k": 4})

    tab1, tab2 = st.tabs(["💬 Document Chat", "📝 Generate Question Paper"])

    with tab1:
        st.subheader("Ask questions from your uploaded files")
        user_query = st.text_input("Enter your question:")
        if user_query:
            with st.spinner("Searching files..."):
                docs = retriever.invoke(user_query)
                context = "\n\n".join([doc.page_content for doc in docs])
                prompt = f"""
                You are a strict academic study assistant. Answer the question based ONLY on the provided context below.
                If the answer cannot be found in the context, explicitly reply: "I cannot answer this based on the uploaded material."

                Context:
                {context}

                Question: {user_query}

                Answer:
                """
                response = llm.invoke(prompt)
                st.write("**Answer:**")
                st.info(response.content)

    with tab2:
        st.subheader("Generate Full Examination Paper")
        col1, col2, col3 = st.columns(3)
        with col1:
            num_mcqs = st.number_input("Number of MCQs", min_value=0, max_value=20, value=5)
        with col2:
            num_short = st.number_input("Short Questions", min_value=0, max_value=10, value=3)
        with col3:
            num_long = st.number_input("Long Questions", min_value=0, max_value=5, value=2)

        paper_title = st.text_input("Subject / Paper Title", value="Midterm Examination")

        if st.button("Generate Paper"):
            with st.spinner("Extracting content and formatting exam paper..."):
                docs = retriever.invoke("Key concepts, definitions, summaries, and main topics")
                context = "\n\n".join([doc.page_content for doc in docs])
                paper_prompt = f"""
                Using ONLY the provided context, generate a complete exam question paper titled '{paper_title}'.
                
                SECTION A: MULTIPLE CHOICE QUESTIONS ({num_mcqs} Questions)
                Create {num_mcqs} MCQs based strictly on the uploaded text with 4 options (A, B, C, D) each.

                SECTION B: SHORT ANSWER QUESTIONS ({num_short} Questions)
                Create {num_short} short-answer questions.

                SECTION C: LONG ANSWER QUESTIONS ({num_long} Questions)
                Create {num_long} essay or analytical long questions.

                Context:
                {context}
                """
                
                response = llm.invoke(paper_prompt)
                st.session_state['generated_paper'] = response.content

        if 'generated_paper' in st.session_state:
            st.markdown("### Preview Generated Paper")
            st.text_area("Paper Preview", st.session_state['generated_paper'], height=350)
            
            pdf_data = create_pdf_paper(st.session_state['generated_paper'])
            
            st.download_button(
                label="📄 Download Exam Paper (PDF)",
                data=pdf_data,
                file_name="Question_Paper.pdf",
                mime="application/pdf"
            )

else:
    st.info("Please enter your Gemini API Key and upload at least one file to start.")
